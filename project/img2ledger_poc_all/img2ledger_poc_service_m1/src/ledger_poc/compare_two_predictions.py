from __future__ import annotations

"""Compare two predicted business.xlsx outputs (engine A vs engine B).

This is used by the Web service "compare mode" to generate a human-readable
Excel diff without requiring any ground-truth.

Matching strategy
-----------------
We primarily match rows by numeric tuple: (量/斤, 单价, 销售额).
This is robust when text fields (客户名称/品名) are noisy.

Thresholds are configurable but default to conservative values.
"""

from dataclasses import dataclass
from typing import Tuple

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows


@dataclass
class MatchConfig:
    jin_tol: float = 0.2
    price_tol: float = 5.0
    amount_tol: float = 50.0


NUM_COLS = [
    "量/只",
    "量/斤",
    "单价",
    "扣款/优惠",
    "销售额",
    "其它费用",
    "应收金额",
    "未收金额",
    "已收金额",
]
TEXT_COLS = ["日期", "单号", "客户名称", "品名", "收款方式", "收款日期", "备注"]


def _norm_num(s):
    return pd.to_numeric(s, errors="coerce")


def _norm_text(s):
    if s is None or (isinstance(s, float) and np.isnan(s)):
        return ""
    t = str(s).replace(" ", "").replace("\n", "")
    return "" if t.lower() == "nan" else t


def _pair_cost(a, b):
    # weighted distance
    dj = abs(a["量/斤"] - b["量/斤"])
    dp = abs(a["单价"] - b["单价"])
    da = abs(a["销售额"] - b["销售额"])
    return dj * 10 + dp * 5 + da * 0.1


def match_rows(df_a: pd.DataFrame, df_b: pd.DataFrame, cfg: MatchConfig = MatchConfig()):
    pairs = []
    for i, ra in df_a.iterrows():
        if pd.isna(ra.get("量/斤")) or pd.isna(ra.get("单价")) or pd.isna(ra.get("销售额")):
            continue
        for j, rb in df_b.iterrows():
            if pd.isna(rb.get("量/斤")) or pd.isna(rb.get("单价")) or pd.isna(rb.get("销售额")):
                continue
            if abs(ra["量/斤"] - rb["量/斤"]) > cfg.jin_tol:
                continue
            if abs(ra["单价"] - rb["单价"]) > cfg.price_tol:
                continue
            if abs(ra["销售额"] - rb["销售额"]) > cfg.amount_tol:
                continue
            pairs.append((_pair_cost(ra, rb), i, j))

    pairs.sort(key=lambda x: x[0])
    used_a, used_b = set(), set()
    matched = []
    for c, i, j in pairs:
        if i in used_a or j in used_b:
            continue
        used_a.add(i)
        used_b.add(j)
        matched.append((i, j, float(c)))

    unmatched_a = [i for i in df_a.index if i not in used_a]
    unmatched_b = [j for j in df_b.index if j not in used_b]
    return matched, unmatched_a, unmatched_b


def build_compare_frames(df_a: pd.DataFrame, df_b: pd.DataFrame, matched):
    rows = []
    for i, j, c in matched:
        a = df_a.loc[i]
        b = df_b.loc[j]
        rows.append(
            {
                "match_cost": c,
                "a_index": int(i),
                "b_index": int(j),
                "客户名称_a": _norm_text(a.get("客户名称")),
                "客户名称_b": _norm_text(b.get("客户名称")),
                "品名_a": _norm_text(a.get("品名")),
                "品名_b": _norm_text(b.get("品名")),
                "量/只_a": a.get("量/只"),
                "量/只_b": b.get("量/只"),
                "量/斤_a": a.get("量/斤"),
                "量/斤_b": b.get("量/斤"),
                "单价_a": a.get("单价"),
                "单价_b": b.get("单价"),
                "销售额_a": a.get("销售额"),
                "销售额_b": b.get("销售额"),
                "收款方式_a": _norm_text(a.get("收款方式")),
                "收款方式_b": _norm_text(b.get("收款方式")),
                "收款日期_a": _norm_text(a.get("收款日期")),
                "收款日期_b": _norm_text(b.get("收款日期")),
            }
        )
    return pd.DataFrame(rows)


def export_compare_xlsx(
    business_a_path: str,
    business_b_path: str,
    out_xlsx_path: str,
    engine_a_name: str = "engine_a",
    engine_b_name: str = "engine_b",
    match_cfg: MatchConfig = MatchConfig(),
):
    df_a = pd.read_excel(business_a_path)
    df_b = pd.read_excel(business_b_path)

    for df in (df_a, df_b):
        for c in NUM_COLS:
            if c in df.columns:
                df[c] = _norm_num(df[c])

    matched, ua, ub = match_rows(df_a, df_b, match_cfg)
    diff = build_compare_frames(df_a, df_b, matched).sort_values("a_index") if matched else pd.DataFrame()

    summary = pd.DataFrame(
        [
            {"metric": f"{engine_a_name}_rows", "value": len(df_a)},
            {"metric": f"{engine_b_name}_rows", "value": len(df_b)},
            {"metric": "matched_rows", "value": len(matched)},
            {"metric": f"unmatched_{engine_a_name}", "value": len(ua)},
            {"metric": f"unmatched_{engine_b_name}", "value": len(ub)},
        ]
    )

    wb = Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    center = Alignment(vertical="center", wrap_text=True)

    def add_sheet(name: str, df: pd.DataFrame):
        ws = wb.create_sheet(name)
        for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start=1):
            ws.append(row)
            if r_idx == 1:
                for c_idx in range(1, len(row) + 1):
                    cell = ws.cell(r_idx, c_idx)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = center
            else:
                for c_idx in range(1, len(row) + 1):
                    ws.cell(r_idx, c_idx).alignment = center
        ws.freeze_panes = "A2"
        for col in ws.columns:
            col_letter = col[0].column_letter
            max_len = 0
            for cell in col[:50]:
                v = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(v))
            ws.column_dimensions[col_letter].width = min(max(10, max_len + 2), 45)

    add_sheet("summary", summary)
    add_sheet("matched_diff", diff if not diff.empty else pd.DataFrame({"note": ["no matched rows"]}))
    add_sheet(f"unmatched_{engine_a_name}", df_a.loc[ua] if ua else pd.DataFrame({"note": ["none"]}))
    add_sheet(f"unmatched_{engine_b_name}", df_b.loc[ub] if ub else pd.DataFrame({"note": ["none"]}))

    wb.save(out_xlsx_path)
